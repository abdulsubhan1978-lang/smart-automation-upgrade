from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()

# Add a title slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Smart Railway Patrol"
subtitle.text = "Target System Performance"

# Add a slide for the KPIs
bullet_slide_layout = prs.slide_layouts[1]
slide2 = prs.slides.add_slide(bullet_slide_layout)
shapes = slide2.shapes
title_shape = shapes.title
title_shape.text = "Target System Performance"

body_shape = shapes.placeholders[1]
tf = body_shape.text_frame
tf.text = "🎯 Detection Accuracy: 96.8%"

p = tf.add_paragraph()
p.text = "🎯 Precision: 97.2%"

p = tf.add_paragraph()
p.text = "🎯 Recall: 96.1%"

p = tf.add_paragraph()
p.text = "🎯 F1-Score: 96.6%"

p = tf.add_paragraph()
p.text = "🎯 Real-Time Speed: 28 FPS"

p = tf.add_paragraph()
p.text = "🎯 Response Time: <0.2 s"

p = tf.add_paragraph()
p.text = "🎯 False Alarm Rate: 2.3%"

# Set font size for all paragraphs
for paragraph in tf.paragraphs:
    paragraph.font.size = Pt(24)

# Save the presentation
output_path = "Smart_Railway_Patrol_Performance.pptx"
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
